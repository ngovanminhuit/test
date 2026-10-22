#!/usr/bin/env python3
"""
infotrade.py - Huấn luyện trọng số W cho 100 node infotrade dự báo BTC 2 phiên tới.

Cài đặt:  pip install -r requirements.txt   (biến môi trường FRED_API_KEY: tùy chọn)
Đặt cùng thư mục với: Infotrade_100_Node_BTC.xlsx

Lệnh:
  python infotrade.py fetch                    # tải dữ liệu miễn phí về ./data
  python infotrade.py train [--start 2018-01-01] [--move 0.05]
  python infotrade.py score [--date YYYY-MM-DD] # chấm điểm 1 ngày bằng W đã huấn luyện

Mẫu huấn luyện: nến ngày t và t+1 (UTC, Binance BTCUSDT) cùng tăng hoặc cùng giảm,
tổng biến động close(t+1)/close(t-1)-1 >= 5% (hoặc <= -5%). Thông tin lấy ở ngày t-1.

Giá trị node = W x M x D. Trong code: x = D*M (có dấu), điểm = sum(W*x).
Huấn luyện: logistic regression có chặn W trong [-1000,1000], phạt L2 kéo W về
giá trị khởi tạo (vì số mẫu ít). Lambda chọn bằng walk-forward theo năm.

Dữ liệu ngoài nguồn tự động: nhập tay vào
  manual_dm.csv   date,node_id,D,M        (node định tính: chấm trực tiếp D, M)
  manual_raw.csv  date,node_id,value      (node định lượng: số thô có dấu theo chiều chuẩn,
                                           script tự quy ra M theo thang trong file Excel)
"""
import argparse, io, math, os, re, sys, time, datetime as dt
import numpy as np
import pandas as pd
import requests
from scipy.optimize import minimize
from openpyxl import load_workbook
import extra

XLSX = "Infotrade_100_Node_BTC.xlsx"
DATA = "data"
MANUAL_DM = "manual_dm.csv"
MANUAL_RAW = "manual_raw.csv"
OUT_W = "weights_trained.csv"
OUT_XLSX = "Infotrade_100_Node_BTC_trained.xlsx"
OUT_SAMPLES = "samples_features.csv"
LEVELS = [0.1, 0.3, 0.5, 0.75, 1.0]
LAMBDAS = [0.003, 0.01, 0.03, 0.1, 0.3, 1.0]
SIGNAL_TH = 0.05  # ngưỡng kết luận trên % tối đa, giống file Excel

# Thang riêng cho node tự động có đơn vị khác bảng Excel
AUTO_TH = {
    18: [10, 20, 50, 100, 200],   # thay đổi WALCL tuần (tỷ USD)
    95: [10, 20, 25, 35, 40],     # Fear&Greed - 50 (tham lam +, sợ hãi -)
}
AUTO_TH.update({k: v for k, v in extra.EXTRA_TH.items() if k not in AUTO_TH})

# ----------------------------------------------------------------- tiện ích
def log(*a):
    print(*a, flush=True)


def parse_num(v):
    """'≥1000' -> 1000, '2k' -> 2000, '+20%' -> 20; chữ -> None."""
    if v is None:
        return None
    t = str(v).strip().replace(",", ".")
    m = re.match(r"^[≥≤<>+\s]*(-?\d+(?:\.\d+)?)\s*(k)?(?=$|[\s%/]|\b)", t, re.I)
    if not m:
        return None
    x = float(m.group(1))
    return x * 1000 if m.group(2) else x


def load_nodes(path=XLSX):
    if not os.path.exists(path):
        sys.exit(f"Không thấy {path} trong thư mục hiện tại.")
    ws = load_workbook(path)["Nodes"]
    nodes = []
    for r in range(6, 106):
        nid = int(ws.cell(r, 1).value)
        th = [parse_num(ws.cell(r, c).value) for c in range(7, 12)]
        if any(v is None for v in th) or any(b <= a for a, b in zip(th, th[1:])):
            th = None  # thang định tính hoặc giảm dần -> chỉ nhập tay D, M
        nodes.append(dict(id=nid, name=ws.cell(r, 3).value,
                          w0=float(ws.cell(r, 12).value or 0), th=th))
    return nodes


def to_dm(x, th):
    """Số thô có dấu -> D*M theo thang 5 bậc (nội suy tuyến tính, tối thiểu 0.01)."""
    if x is None or th is None:
        return 0.0
    try:
        x = float(x)
    except (TypeError, ValueError):
        return 0.0
    if not math.isfinite(x) or x == 0:
        return 0.0
    m = float(np.interp(abs(x), [0.0] + list(th), [0.0] + LEVELS))
    return math.copysign(max(m, 0.01), x)


def series(key):
    p = os.path.join(DATA, f"{key}.csv")
    if not os.path.exists(p):
        return None
    df = pd.read_csv(p)
    if df.empty:
        return None
    idx = pd.to_datetime(df["date"], utc=True).dt.tz_convert(None).dt.normalize()
    s = pd.Series(pd.to_numeric(df["value"], errors="coerce").values, index=idx)
    s = s[~s.index.duplicated(keep="last")].sort_index().dropna()
    return s if len(s) else None


def load_btc():
    p = os.path.join(DATA, "btc.csv")
    if not os.path.exists(p):
        sys.exit("Chưa có data/btc.csv - chạy: python infotrade.py fetch")
    df = pd.read_csv(p)
    df.index = pd.to_datetime(df["date"]).dt.normalize()
    return df[["open", "high", "low", "close", "volume"]].astype(float).sort_index()


# ----------------------------------------------------------------- tải dữ liệu
SES = requests.Session()
SES.headers["User-Agent"] = "Mozilla/5.0 infotrade"


def get(url, timeout=30, tries=3, **kw):
    for i in range(tries):
        try:
            r = SES.get(url, timeout=timeout, **kw)
            r.raise_for_status()
            return r
        except requests.HTTPError as e:
            code = e.response.status_code if e.response is not None else 0
            if code in (400, 401, 403, 404, 451) or i == tries - 1:
                raise  # lỗi chặn/không tồn tại: không thử lại
            time.sleep(3 * (i + 1))
        except Exception:
            if i == tries - 1:
                raise
            time.sleep(3 * (i + 1))


def save(df, key):
    os.makedirs(DATA, exist_ok=True)
    df.to_csv(os.path.join(DATA, f"{key}.csv"), index=False)
    log(f"  {key}: {len(df)} dòng")


def ms_to_date(ms):
    return pd.to_datetime(ms, unit="ms").dt.strftime("%Y-%m-%d")


def _klines(base):
    start, rows = 1502928000000, []
    while True:
        js = get(f"{base}/api/v3/klines", params=dict(symbol="BTCUSDT", interval="1d",
                                                     startTime=start, limit=1000)).json()
        if not js:
            break
        rows += js
        start = js[-1][0] + 86_400_000
        if len(js) < 1000:
            break
    df = pd.DataFrame([r[:6] for r in rows],
                      columns=["t", "open", "high", "low", "close", "volume"])
    df["date"] = ms_to_date(df["t"])
    return df


def _coinbase_ohlcv():
    rows, start, now = [], dt.datetime(2017, 8, 1), dt.datetime.utcnow()
    while start < now:
        end = min(start + dt.timedelta(days=299), now)
        js = get("https://api.exchange.coinbase.com/products/BTC-USD/candles",
                 params=dict(granularity=86400, start=start.isoformat(), end=end.isoformat())).json()
        rows += js if isinstance(js, list) else []
        start = end + dt.timedelta(days=1)
        time.sleep(0.4)
    df = pd.DataFrame(rows, columns=["t", "low", "high", "open", "close", "volume"])
    df["date"] = pd.to_datetime(df["t"], unit="s").dt.strftime("%Y-%m-%d")
    return df.drop_duplicates("date").sort_values("date")


def fetch_btc():
    df = None
    for base in ("https://api.binance.com", "https://data-api.binance.vision"):
        try:
            df = _klines(base)
            log(f"  nguồn: {base}")
            break
        except Exception as e:
            log(f"  ! {base}: {e}")
    if df is None or df.empty:
        df = _coinbase_ohlcv()
        log("  nguồn: Coinbase BTC-USD")
    today = dt.datetime.now(dt.timezone.utc).strftime("%Y-%m-%d")
    df = df[df["date"] < today]  # bỏ nến chưa đóng
    save(df[["date", "open", "high", "low", "close", "volume"]], "btc")


def fetch_funding():
    try:
        url, start, rows = "https://fapi.binance.com/fapi/v1/fundingRate", 1568102400000, []
        while True:
            js = get(url, params=dict(symbol="BTCUSDT", startTime=start, limit=1000)).json()
            if not js:
                break
            rows += js
            start = js[-1]["fundingTime"] + 1
            if len(js) < 1000:
                break
        df = pd.DataFrame(rows)
        df["date"] = ms_to_date(df["fundingTime"])
        df["value"] = df["fundingRate"].astype(float) * 100  # %/8h
        log("  nguồn: Binance")
    except Exception as e:
        log(f"  ! Binance: {e} -> dùng Deribit BTC-PERPETUAL")
        rows, start, now = [], dt.datetime(2019, 5, 1), dt.datetime.utcnow()
        while start < now:
            end = min(start + dt.timedelta(days=30), now)
            js = get("https://www.deribit.com/api/v2/public/get_funding_rate_history",
                     params=dict(instrument_name="BTC-PERPETUAL",
                                 start_timestamp=int(start.timestamp() * 1000),
                                 end_timestamp=int(end.timestamp() * 1000))).json()
            rows += js.get("result", [])
            start = end
            time.sleep(0.2)
        df = pd.DataFrame(rows)
        df["date"] = ms_to_date(df["timestamp"])
        df["value"] = df["interest_8h"].astype(float) * 100  # %/8h
    save(df.groupby("date", as_index=False)["value"].mean(), "funding")


YF = {"spx": "^GSPC", "ndx": "^IXIC", "vix": "^VIX", "dxy": "DX-Y.NYB", "gold": "GC=F",
      "oil": "CL=F", "usdjpy": "JPY=X", "mstr": "MSTR", "coin": "COIN", "eth": "ETH-USD"}


def fetch_yf():
    import yfinance as yf
    for key, tk in YF.items():
        try:
            df = yf.download(tk, start="2017-01-01", progress=False, auto_adjust=False)
            if isinstance(df.columns, pd.MultiIndex):
                df.columns = df.columns.get_level_values(0)
            out = pd.DataFrame({"date": pd.to_datetime(df.index).strftime("%Y-%m-%d"),
                                "value": df["Close"].values})
            save(out, key)
        except Exception as e:
            log(f"  ! {key} ({tk}) lỗi: {e}")


FRED = ["DGS10", "DGS2", "WALCL", "WTREGEN", "RRPONTSYD"]


YF_FALLBACK = {"DGS10": ["^TNX"], "DGS2": ["2YY=F", "^FVX"]}  # khi FRED lỗi


def fred_csv(sid):
    r = get("https://fred.stlouisfed.org/graph/fredgraph.csv", timeout=60, tries=2,
            params={"id": sid, "cosd": "2015-01-01"},
            headers={"User-Agent": "Mozilla/5.0 (X11; Linux x86_64) AppleWebKit/537.36 "
                                   "(KHTML, like Gecko) Chrome/126 Safari/537.36"})
    df = pd.read_csv(io.StringIO(r.text))
    df.columns = ["date", "value"]
    df["value"] = pd.to_numeric(df["value"], errors="coerce")
    return df.dropna()


def fetch_fred():
    fails = 0
    for sid in FRED:
        if fails >= 2:
            log(f"  ! FRED không phản hồi - bỏ qua {sid}")
        else:
            try:
                save(fred_csv(sid), sid)
                continue
            except Exception as e:
                fails += 1
                log(f"  ! FRED {sid}: {type(e).__name__}")
        for tk in YF_FALLBACK.get(sid, []):
            try:
                import yfinance as yf
                df = yf.download(tk, start="2017-01-01", progress=False, auto_adjust=False)
                if isinstance(df.columns, pd.MultiIndex):
                    df.columns = df.columns.get_level_values(0)
                if len(df):
                    save(pd.DataFrame({"date": pd.to_datetime(df.index).strftime("%Y-%m-%d"),
                                       "value": df["Close"].values}).dropna(), sid)
                    log(f"    (dự phòng Yahoo {tk})")
                    break
            except Exception as e:
                log(f"  ! Yahoo {tk}: {e}")


def fetch_fng():
    js = get("https://api.alternative.me/fng/", params={"limit": 0, "format": "json"}).json()["data"]
    df = pd.DataFrame(js)
    df["date"] = pd.to_datetime(df["timestamp"].astype(int), unit="s").dt.strftime("%Y-%m-%d")
    df["value"] = df["value"].astype(float)
    save(df[["date", "value"]], "fng")


def fetch_stable():
    for key, sid in {"usdt": 1, "usdc": 2}.items():
        try:
            js = get("https://stablecoins.llama.fi/stablecoincharts/all",
                     params={"stablecoin": sid}).json()
            rows = []
            for x in js:
                tot = x.get("totalCirculatingUSD") or x.get("totalCirculating") or {}
                if "peggedUSD" in tot:
                    d = dt.datetime.fromtimestamp(int(x["date"]), dt.timezone.utc)
                    rows.append((d.strftime("%Y-%m-%d"), float(tot["peggedUSD"])))
            save(pd.DataFrame(rows, columns=["date", "value"]), key)
        except Exception as e:
            log(f"  ! {key} lỗi: {e}")


def fetch_mempool():
    js = get("https://mempool.space/api/v1/mining/hashrate/all").json()
    hr = pd.DataFrame(js.get("hashrates", []))
    hr["date"] = pd.to_datetime(hr["timestamp"], unit="s").dt.strftime("%Y-%m-%d")
    save(hr.rename(columns={"avgHashrate": "value"})[["date", "value"]], "hashrate")
    df = pd.DataFrame(js.get("difficulty", []))
    if not df.empty:
        tcol = "time" if "time" in df.columns else "timestamp"
        df["date"] = pd.to_datetime(df[tcol], unit="s").dt.strftime("%Y-%m-%d")
        save(df.rename(columns={"difficulty": "value"})[["date", "value"]], "difficulty")


def cmd_fetch(_):
    for name, fn in [("BTC nến ngày (Binance)", fetch_btc), ("Funding (Binance)", fetch_funding),
                     ("Yahoo Finance", fetch_yf), ("FRED", fetch_fred),
                     ("Fear & Greed", fetch_fng), ("Stablecoin (DefiLlama)", fetch_stable),
                     ("Hashrate (mempool.space)", fetch_mempool)]:
        log(f"Tải {name} ...")
        try:
            fn()
        except ImportError:
            log("  ! Thiếu thư viện: pip install yfinance")
        except Exception as e:
            log(f"  ! Lỗi: {e}")
    extra.fetch_all(get, save, log)
    log("Xong. Dữ liệu ở thư mục ./data")


# ----------------------------------------------------------------- dựng đặc trưng
def tech_signal(btc):
    """Node 99: phá mốc kỹ thuật, trả trực tiếp D*M."""
    c = btc["close"]
    p = c.shift(1)
    ma50, ma200 = c.rolling(50).mean(), c.rolling(200).mean()
    hi30, lo30 = c.rolling(30).max().shift(1), c.rolling(30).min().shift(1)
    ath = c.cummax().shift(1)
    up = lambda ma: (c > ma) & (p <= ma.shift(1))
    dn = lambda ma: (c < ma) & (p >= ma.shift(1))
    conds = [c > ath, c > hi30, up(ma200), up(ma50), c < lo30, dn(ma200), dn(ma50)]
    vals = [1.0, 0.75, 0.5, 0.3, -0.75, -0.5, -0.3]
    return pd.Series(np.select(conds, vals, 0.0), index=c.index)


def raw_auto(btc):
    """Số thô có dấu theo 'chiều chuẩn' cho các node lấy tự động."""
    raw = {}

    def pc(key):
        s = series(key)
        return None if s is None else s.pct_change() * 100

    s = series("DGS10")
    if s is not None:
        raw[15] = s.diff() * 100                       # bp, 10Y tăng
    s = series("DGS2")
    if s is not None:
        raw[16] = -s.diff() * 100                      # 2Y giảm ~ bull steepening
    raw[17] = pc("dxy")
    walcl, tga, rrp = series("WALCL"), series("WTREGEN"), series("RRPONTSYD")
    if walcl is not None:
        raw[18] = walcl.diff() / 1000                  # tỷ USD/tuần
        if tga is not None and rrp is not None:
            full = pd.date_range(min(walcl.index[0], tga.index[0], rrp.index[0]),
                                 max(walcl.index[-1], tga.index[-1], rrp.index[-1]))
            t = tga.reindex(full).ffill()
            r = rrp.reindex(full).ffill().fillna(0)
            net = walcl / 1000 - t.reindex(walcl.index) / 1000 - r.reindex(walcl.index)
            raw[19] = net.diff()                       # tỷ USD/tuần
    for nid, key in [(21, "spx"), (22, "ndx"), (23, "vix"), (24, "gold"), (25, "oil"), (92, "eth")]:
        raw[nid] = pc(key)
    parts = [x for x in (pc("mstr"), pc("coin")) if x is not None]
    if parts:
        raw[27] = pd.concat(parts, axis=1).mean(axis=1)
    j = pc("usdjpy")
    if j is not None:
        raw[28] = -j                                   # JPY tăng = USD/JPY giảm
    for nid, key in [(39, "usdt"), (40, "usdc")]:
        s = series(key)
        if s is not None:
            raw[nid] = s.diff() / 1e9                  # tỷ USD/ngày
    raw[45] = series("funding")
    v, ret = btc["volume"], btc["close"].pct_change()
    ratio = v / v.rolling(20).mean().shift(1) - 1
    raw[56] = pd.Series(np.where(ratio > 0, np.sign(ret) * ratio * 100, 0.0), index=btc.index)
    hr = series("hashrate")
    if hr is not None:
        raw[62] = hr.resample("D").mean().interpolate().pct_change(7) * 100
    dif = series("difficulty")
    if dif is not None:
        raw[63] = dif.pct_change() * 100
    fg = series("fng")
    if fg is not None:
        raw[95] = fg - 50
    return {k: v for k, v in raw.items() if v is not None}


def build_features(nodes, btc, idx):
    byid = {n["id"]: n for n in nodes}
    X = pd.DataFrame(0.0, index=idx, columns=[n["id"] for n in nodes])
    raw = raw_auto(btc)
    raw2, direct = extra.raw_extra(series, btc)
    for nid, s in raw2.items():
        raw.setdefault(nid, s)
    for nid, s in raw.items():
        th = AUTO_TH.get(nid) or byid[nid]["th"]
        X[nid] = [to_dm(x, th) for x in s.reindex(idx).values]
    X[99] = tech_signal(btc).reindex(idx).fillna(0.0)
    for nid, s in direct.items():
        X[nid] = s.reindex(idx).fillna(0.0).clip(-1, 1).values
    auto = sorted(set(raw) | set(direct) | {99})

    if os.path.exists(MANUAL_RAW):
        df = pd.read_csv(MANUAL_RAW, comment="#")
        for _, r in df.iterrows():
            d, nid = pd.Timestamp(r["date"]).normalize(), int(r["node_id"])
            if d not in X.index or nid not in byid:
                continue
            th = AUTO_TH.get(nid) or byid[nid]["th"]
            if th is None:
                log(f"  ! manual_raw: node {nid} là thang định tính, dùng manual_dm.csv")
                continue
            X.loc[d, nid] = to_dm(r["value"], th)
    if os.path.exists(MANUAL_DM):
        df = pd.read_csv(MANUAL_DM, comment="#")
        for _, r in df.iterrows():
            d, nid = pd.Timestamp(r["date"]).normalize(), int(r["node_id"])
            D, M = int(r["D"]), float(r["M"])
            if d not in X.index or nid not in byid:
                continue
            if D not in (-1, 0, 1) or not 0.01 <= M <= 1:
                log(f"  ! manual_dm: bỏ dòng sai {r['date']} node {nid} D={D} M={M}")
                continue
            X.loc[d, nid] = D * M
    return X, auto


def make_samples(btc, start, move):
    c = btc["close"]
    r1 = c.pct_change()
    r2 = r1.shift(-1)
    tot = c.shift(-1) / c.shift(1) - 1
    up = (r1 > 0) & (r2 > 0) & (tot >= move)
    dn = (r1 < 0) & (r2 < 0) & (tot <= -move)
    y = pd.Series(np.where(up, 1, np.where(dn, -1, 0)), index=c.index)
    y = y[(y != 0) & (y.index >= start)]
    return pd.DataFrame({"candle1": y.index, "info_date": y.index - pd.Timedelta(days=1),
                         "y": y.values, "move": tot.reindex(y.index).values})


# ----------------------------------------------------------------- huấn luyện
def class_weights(y):
    w = np.ones_like(y, dtype=float)
    for c in (-1, 1):
        k = (y == c).sum()
        if k:
            w[y == c] = len(y) / (2 * k)
    return w


def fit(X, y, w0, lam):
    sw = class_weights(y)
    tot = sw.sum()

    def f(w):
        z = -y * (X @ w)
        loss = np.logaddexp(0, z)
        p = 0.5 * (1 + np.tanh(z / 2))  # sigmoid(z), ổn định số
        L = (sw * loss).sum() / tot + lam * ((w - w0) ** 2).sum()
        g = X.T @ (sw * (-y) * p) / tot + 2 * lam * (w - w0)
        return L, g

    res = minimize(f, w0.copy(), jac=True, method="L-BFGS-B", bounds=[(-1, 1)] * len(w0))
    return res.x


def accuracy(X, y, w):
    pred = np.sign(X @ w)
    return int((pred == y).sum())


def cmd_train(a):
    nodes = load_nodes()
    ids = [n["id"] for n in nodes]
    w0 = np.array([n["w0"] for n in nodes]) / 1000
    btc = load_btc()
    X_all, auto = build_features(nodes, btc, btc.index)
    S = make_samples(btc, pd.Timestamp(a.start), a.move)
    S = S[S["info_date"].isin(X_all.index)].reset_index(drop=True)
    if len(S) < 20:
        sys.exit(f"Chỉ có {len(S)} mẫu - quá ít để huấn luyện.")
    X = X_all.loc[S["info_date"].values, ids].values
    y = S["y"].values.astype(float)
    years = S["candle1"].dt.year.values
    log(f"Mẫu: {len(S)} (tăng {int((y > 0).sum())}, giảm {int((y < 0).sum())}), "
        f"{S['candle1'].min():%Y-%m-%d} -> {S['candle1'].max():%Y-%m-%d}")
    has = (np.abs(X) > 0).any(0)
    w0 = np.where(has, w0, 0.0)  # node chưa có dữ liệu: W = 0
    log(f"Node tự động: {len(auto)} | node có dữ liệu trong mẫu: {int(has.sum())}/100 "
        f"(còn lại W = 0)")

    test_years = [Y for Y in sorted(set(years)) if (years < Y).sum() >= 30]
    best_lam, rows = 0.1, []
    if test_years:
        log("\nWalk-forward (huấn luyện các năm trước, kiểm tra năm sau):")
        scores = {}
        for lam in LAMBDAS:
            hit = n = 0
            for Y in test_years:
                tr, te = years < Y, years == Y
                w = fit(X[tr], y[tr], w0, lam)
                hit += accuracy(X[te], y[te], w)
                n += te.sum()
            scores[lam] = hit / n
            log(f"  lambda={lam:<6} độ chính xác={scores[lam]:.1%}")
        best_lam = max(LAMBDAS, key=lambda l: (round(scores[l], 4), l))
        log(f"  -> chọn lambda={best_lam}")
        log("\n  Năm   Mẫu  Đúng(W học)  Đúng(W khởi tạo)  Đoán theo đa số")
        for Y in test_years:
            tr, te = years < Y, years == Y
            w = fit(X[tr], y[tr], w0, best_lam)
            maj = 1.0 if (y[tr] > 0).sum() >= (y[tr] < 0).sum() else -1.0
            n = int(te.sum())
            rows.append((Y, n, accuracy(X[te], y[te], w), accuracy(X[te], y[te], w0),
                         int((y[te] == maj).sum())))
            r = rows[-1]
            log(f"  {Y}  {n:>4}  {r[2] / n:>10.1%}  {r[3] / n:>15.1%}  {r[4] / n:>14.1%}")
        tn = sum(r[1] for r in rows)
        log(f"  Tổng  {tn:>4}  {sum(r[2] for r in rows) / tn:>10.1%}  "
            f"{sum(r[3] for r in rows) / tn:>15.1%}  {sum(r[4] for r in rows) / tn:>14.1%}")
    else:
        log("Không đủ năm để walk-forward, dùng lambda=0.1")

    w = fit(X, y, w0, best_lam)
    W = np.round(w * 1000).astype(int)
    active = (np.abs(X) > 0).sum(0)
    out = pd.DataFrame({"id": ids, "name": [n["name"] for n in nodes],
                        "W_init": (w0 * 1000).astype(int), "W_trained": W,
                        "so_mau_co_tin": active,
                        "nguon": ["tu dong" if i in auto else ("tay" if active[k] else "chua co")
                                  for k, i in enumerate(ids)]})
    out.to_csv(OUT_W, index=False, encoding="utf-8-sig")

    feat = pd.DataFrame(X, columns=[f"n{i}" for i in ids])
    S2 = pd.concat([S, feat, pd.Series(X @ W, name="score")], axis=1)
    S2.to_csv(OUT_SAMPLES, index=False, encoding="utf-8-sig")

    wb = load_workbook(XLSX)
    ws = wb["Nodes"]
    for k in range(100):
        ws.cell(6 + k, 12).value = int(W[k])
    ws.cell(5, 12).value = "Trọng số W (đã huấn luyện)"
    wb.save(OUT_XLSX)

    chg = out.assign(d=(out.W_trained - out.W_init).abs()).sort_values("d", ascending=False)
    log("\nNode thay đổi W nhiều nhất:")
    for _, r in chg.head(15).iterrows():
        log(f"  {r.id:>3} {r['name'][:38]:<38} {r.W_init:>5} -> {r.W_trained:>5}  ({r.so_mau_co_tin} mẫu)")
    log(f"\nĐã ghi: {OUT_W}, {OUT_XLSX}, {OUT_SAMPLES}")


# ----------------------------------------------------------------- chấm điểm
def cmd_score(a):
    nodes = load_nodes()
    ids = [n["id"] for n in nodes]
    name = {n["id"]: n["name"] for n in nodes}
    if os.path.exists(OUT_W):
        W = pd.read_csv(OUT_W).set_index("id").loc[ids, "W_trained"].values.astype(float)
        src = OUT_W
    else:
        W = np.array([n["w0"] for n in nodes])
        src = "W khởi tạo (chưa huấn luyện)"
    btc = load_btc()
    d = pd.Timestamp(a.date).normalize() if a.date else btc.index.max()
    idx = pd.date_range(btc.index.min(), max(d, btc.index.max()))
    X_all, _ = build_features(nodes, btc, idx)
    x = X_all.loc[d, ids].values
    contrib = W * x
    total, mx = contrib.sum(), np.abs(W).sum()
    pct = total / mx if mx else 0
    prob = 1 / (1 + math.exp(-float(x @ (W / 1000))))
    log(f"Ngày thông tin: {d:%Y-%m-%d}  -> dự báo 2 phiên {d + pd.Timedelta(days=1):%d/%m} "
        f"và {d + pd.Timedelta(days=2):%d/%m}   (W: {src})\n")
    log(f"  {'ID':>3}  {'Node':<40} {'W':>6} {'D*M':>6} {'Giá trị':>9}")
    for k in np.argsort(-np.abs(contrib)):
        if x[k] != 0:
            log(f"  {ids[k]:>3}  {name[ids[k]][:40]:<40} {int(W[k]):>6} {x[k]:>6.2f} {contrib[k]:>9.1f}")
    kl = "Nghiêng TĂNG" if pct >= SIGNAL_TH else ("Nghiêng GIẢM" if pct <= -SIGNAL_TH else "Trung tính")
    log(f"\nTổng điểm: {total:,.1f} / tối đa {mx:,.0f}  = {pct:.1%}")
    log(f"Xác suất tăng (mô hình): {prob:.0%}   Kết luận: {kl}")
    log(f"Node có tin: {int((x != 0).sum())}/100 - nhớ nhập tin định tính vào {MANUAL_DM}")


# ----------------------------------------------------------------- giả lập
def max_drawdown(eq):
    peak = np.maximum.accumulate(eq)
    return float((eq / peak - 1).min())


def cmd_simulate(a):
    """Giả lập ngoài mẫu: năm Y dùng W học từ các năm trước, chấm điểm MỌI ngày của năm Y.
    Tín hiệu ở ngày d (sau khi đóng nến d) -> vào lệnh giá đóng cửa d, giữ 2 phiên, đóng ở close d+2.
    Lệnh không chồng nhau (đang giữ lệnh thì bỏ qua tín hiệu mới). Phí 0.1% mỗi chiều."""
    nodes = load_nodes()
    ids = [n["id"] for n in nodes]
    w0 = np.array([n["w0"] for n in nodes]) / 1000
    btc = load_btc()
    X_all, _ = build_features(nodes, btc, btc.index)
    S = make_samples(btc, pd.Timestamp(a.start), a.move)
    S = S[S["info_date"].isin(X_all.index)].reset_index(drop=True)
    Xs = X_all.loc[S["info_date"].values, ids].values
    y = S["y"].values.astype(float)
    years = S["candle1"].dt.year.values
    has = (np.abs(Xs) > 0).any(0)
    w0 = np.where(has, w0, 0.0)
    lam = a.lam
    c = btc["close"]
    fwd = c.shift(-2) / c - 1
    test_years = [Y for Y in sorted(set(years)) if (years < Y).sum() >= 30]
    parts = []
    for Y in test_years:
        w = fit(Xs[years < Y], y[years < Y], w0, lam)
        days = btc.index[(btc.index.year == Y) & fwd.notna().values]
        z = X_all.loc[days, ids].values @ w
        parts.append(pd.DataFrame({"date": days, "p_tang": 1 / (1 + np.exp(-z)),
                                   "ret2": fwd.reindex(days).values}))
    D = pd.concat(parts, ignore_index=True)
    log(f"Giả lập {D.date.min():%Y-%m-%d} -> {D.date.max():%Y-%m-%d} | {len(D)} ngày | "
        f"node dùng: {int(has.sum())}/100 | lambda={lam} | phí {a.fee:.2%}/chiều\n")

    bh = float(c.reindex(D.date).iloc[-1] / c.reindex(D.date).iloc[0])
    log(f"Mua & giữ BTC cùng kỳ: {bh - 1:+.1%}\n")
    log(f"  {'Ngưỡng':<8}{'Lệnh':>6}{'Long':>6}{'Short':>6}{'Thắng':>8}{'TB/lệnh':>9}"
        f"{'Tổng lãi':>11}{'Sụt giảm max':>14}")
    best = None
    for th in (0.52, 0.55, 0.60, 0.65, 0.70):
        trades, skip = [], -1
        for i, r in D.iterrows():
            if i <= skip:
                continue
            side = 1 if r.p_tang >= th else (-1 if r.p_tang <= 1 - th else 0)
            if side:
                trades.append((r.date, side, r.p_tang, side * r.ret2 - 2 * a.fee))
                skip = i + 1
        T = pd.DataFrame(trades, columns=["date", "side", "p_tang", "ret"])
        if T.empty:
            log(f"  {th:<8.2f}{0:>6}")
            continue
        eq = np.cumprod(1 + T["ret"].values)
        log(f"  {th:<8.2f}{len(T):>6}{int((T.side > 0).sum()):>6}{int((T.side < 0).sum()):>6}"
            f"{(T.ret > 0).mean():>8.1%}{T.ret.mean():>+9.2%}{eq[-1] - 1:>+11.1%}{max_drawdown(eq):>14.1%}")
        if th == a.th:
            best = T.assign(von=eq)
    if best is not None:
        os.makedirs("results", exist_ok=True)
        best.to_csv("results/simulation_trades.csv", index=False, encoding="utf-8-sig")
        best["nam"] = pd.to_datetime(best.date).dt.year
        log(f"\nTheo năm (ngưỡng {a.th}):")
        for Y, g in best.groupby("nam"):
            yr = c[c.index.year == Y]
            log(f"  {Y}: {len(g):>3} lệnh, thắng {(g.ret > 0).mean():.0%}, "
                f"lãi {np.prod(1 + g.ret) - 1:+.1%} | BTC cả năm {yr.iloc[-1] / yr.iloc[0] - 1:+.1%}")
    D.to_csv("results/simulation_daily.csv", index=False, encoding="utf-8-sig")
    log("\nĐã ghi: results/simulation_trades.csv, results/simulation_daily.csv")


def main():
    try:
        sys.stdout.reconfigure(encoding="utf-8")  # console Windows hiển thị tiếng Việt
    except Exception:
        pass
    ap =argparse.ArgumentParser(description="Huấn luyện node infotrade BTC")
    sub = ap.add_subparsers(dest="cmd", required=True)
    sub.add_parser("fetch")
    t = sub.add_parser("train")
    t.add_argument("--start", default="2018-01-01")
    t.add_argument("--move", type=float, default=0.05, help="tổng biến động 2 phiên (0.05 = 5%%)")
    m = sub.add_parser("simulate")
    m.add_argument("--start", default="2018-01-01")
    m.add_argument("--move", type=float, default=0.05)
    m.add_argument("--lam", type=float, default=0.1)
    m.add_argument("--th", type=float, default=0.55, help="ngưỡng xác suất vào lệnh")
    m.add_argument("--fee", type=float, default=0.001)
    s = sub.add_parser("score")
    s.add_argument("--date")
    a = ap.parse_args()
    {"fetch": cmd_fetch, "train": cmd_train, "score": cmd_score, "simulate": cmd_simulate}[a.cmd](a)


if __name__ == "__main__":
    main()
